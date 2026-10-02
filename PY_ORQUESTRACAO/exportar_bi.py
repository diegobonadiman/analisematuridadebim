# -*- coding: utf-8 -*-
"""Exporta metricas_*.json em CSVs para Power BI.

Gera 4 arquivos em _BI_EXPORT\<AAAA-MM>\:
  bi_modelos.csv   - dimensao (com tamanho_mb, qtd_familias, qtd_instancias)
  bi_checks.csv    - dimensao de checks
  bi_metricas.csv  - fato (modelo x check)
  bi_resumo.csv    - fato agregado por modelo
"""
import os
import re
import sys
import csv
import glob
import json
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from bim_paths import PASTA_SAIDA, PLANILHA_AUDIT

# pasta raiz do BI + subpasta por periodo (para historico)
PASTA_BI_RAIZ = os.path.join(os.path.dirname(PASTA_SAIDA), "_BI_EXPORT")
PERIODO = datetime.now().strftime("%Y-%m")
PASTA_BI = os.path.join(PASTA_BI_RAIZ, PERIODO)

# Checks mensuraveis (mesma lista do consolidar.py)
MAPA_API = {
    5:  "nomenclatura_arquivo",
    6:  "elementos_voados",
    8:  "dwg_refs",
    9:  "modelagem_local",
    10: "nivel_vista_fora_padrao",
    13: "folhas_fora_padrao",
    15: "carimbo_incompleto",
    16: "familias_generico",
    18: "tabelas_fora_padrao",
    23: "total_tags",
    25: "avisos_graves",
    75: "qtd_vinculos_rvt",
    76: "qtd_vinculos_cad",
    77: "qtd_vistas_sem_template",
}

CHECKS_BINARIOS     = {5}
CHECKS_MAIOR_MELHOR = {23}

# Campos do JSON que viram colunas de dimensao em bi_modelos
CAMPOS_DIMENSAO = [
    ("tamanho_mb",             "tamanho_arquivo_mb"),
    ("qtd_familias_unicas",    "qtd_familias_unicas"),
    ("qtd_instancias_familia", "qtd_instancias_familia"),
    ("qtd_worksets",           "qtd_worksets"),
]

RX_PROJETO = re.compile(
    r"^(?P<escritorio>WD|ME|PA|AG)?[-_]?(?P<projeto>\d+|[A-Z]+)[-_]"
    r"(?P<nome>[^_]+)[-_]"
    r"(?P<disciplina>[A-Z]+)[-_]"
    r"(?P<tipo>EX|EP|ARQ)[-_]?"
    r"(?P<revisao>R\d+)",
    re.IGNORECASE
)


# ============================================================
#  HELPERS
# ============================================================
def extrair_dimensoes(nome_arquivo):
    d = {"escritorio": "", "projeto": "", "nome": "",
         "disciplina": "", "tipo": "", "revisao": "", "formato": ""}
    upper = nome_arquivo.upper()
    if ".IFC.RVT" in upper:
        d["formato"] = "IFC_RVT"
    elif ".RVT" in upper:
        d["formato"] = "RVT"
    base = nome_arquivo
    for ext in (".ifc.RVT", ".ifc.rvt", ".RVT", ".rvt"):
        if base.endswith(ext):
            base = base[: -len(ext)]
            break
    m = RX_PROJETO.match(base)
    if m:
        for k in ("escritorio", "projeto", "nome", "disciplina", "tipo", "revisao"):
            v = m.group(k) or ""
            d[k] = v.upper() if k in ("disciplina", "tipo", "revisao") else v
    return d


def carregar_json():
    dados = {}
    for jf in glob.glob(os.path.join(PASTA_SAIDA, "metricas_*.json")):
        try:
            with open(jf, encoding="utf-8") as f:
                dados.update(json.load(f))
        except Exception as e:
            print(u"[AVISO] falha lendo {}: {}".format(jf, e))
    return dados


def carregar_pesos():
    import openpyxl
    wb = openpyxl.load_workbook(PLANILHA_AUDIT, data_only=True)
    ws = wb["ARC"]
    pesos, descricoes, categorias = {}, {}, {}
    for row in MAPA_API:
        p = ws.cell(row=row, column=5).value
        d = ws.cell(row=row, column=2).value or ""
        c = ws.cell(row=row, column=1).value or ""
        pesos[row] = int(p) if isinstance(p, (int, float)) else 0
        descricoes[row] = str(d).strip()
        categorias[row] = str(c).strip()
    return pesos, descricoes, categorias


def calcular_escala(dados):
    escala = {}
    for row in MAPA_API:
        chave = MAPA_API[row]
        valores = [m.get(chave) for m in dados.values()
                   if isinstance(m.get(chave), (int, float)) and m.get(chave) >= 0]
        if valores:
            escala[row] = {"min": min(valores), "max": max(valores),
                           "valores": sorted(valores)}
    return escala


def metrica_para_nota(row, valor, escala):
    if valor is None or valor < 0 or row not in escala:
        return None
    e = escala[row]
    vmin, vmax = e["min"], e["max"]
    if row in CHECKS_BINARIOS:
        return 10 if valor else 0
    if row in CHECKS_MAIOR_MELHOR:
        if valor == 0:
            return 0
        nz = [x for x in e["valores"] if x > 0]
        if not nz:
            return 0
        vmin_nz, vmax_nz = min(nz), max(nz)
        if vmax_nz == vmin_nz:
            return 10
        return round((valor - vmin_nz) / float(vmax_nz - vmin_nz) * 10, 2)
    if vmax == vmin:
        return 10
    return round((1.0 - (valor - vmin) / float(vmax - vmin)) * 10, 2)


def _valor_csv(v):
    """Converte float em string com virgula decimal (locale PT-BR).
    Ints, strings e None passam sem alteracao."""
    if isinstance(v, float):
        s = ("%.4f" % v).rstrip("0").rstrip(".")
        return s.replace(".", ",")
    return v


def escrever_csv(caminho, header, linhas):
    """Grava CSV com ';' como separador e ',' como decimal.
    Power BI e Excel PT-BR leem corretamente."""
    with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        w.writerow(header)
        for l in linhas:
            w.writerow([_valor_csv(v) for v in l])


# ============================================================
#  MAIN
# ============================================================
def main():
    if not os.path.isdir(PASTA_BI):
        os.makedirs(PASTA_BI)

    dados = carregar_json()
    if not dados:
        print(u"[ERRO] Nenhum metricas_*.json em {}".format(PASTA_SAIDA))
        return

    pesos, descricoes, categorias = carregar_pesos()
    escala = calcular_escala(dados)
    hoje = datetime.now().strftime("%Y-%m-%d %H:%M")

    print(u"Modelos carregados: {}".format(len(dados)))
    print(u"Saida: {}".format(PASTA_BI))
    print(u"")

    # -------- bi_modelos.csv --------
    modelos_linhas = []
    mapa_modelo_id = {}
    for i, arq in enumerate(sorted(dados.keys()), 1):
        dim = extrair_dimensoes(arq)
        modelo_id = "M{:03d}".format(i)
        mapa_modelo_id[arq] = modelo_id
        linha = [
            modelo_id,
            arq,
            dim["nome"],
            dim["escritorio"] or "N/A",
            dim["projeto"],
            dim["disciplina"],
            dim["tipo"],
            dim["revisao"],
            dim["formato"],
            hoje,
        ]
        # adiciona colunas de dimensao extra
        for _, chave in CAMPOS_DIMENSAO:
            linha.append(dados[arq].get(chave, -1))
        modelos_linhas.append(linha)

    header_modelos = [
        "modelo_id", "nome_arquivo", "nome_projeto", "escritorio",
        "projeto", "disciplina", "tipo", "revisao", "formato", "data_analise",
    ] + [nome for nome, _ in CAMPOS_DIMENSAO]

    escrever_csv(os.path.join(PASTA_BI, "bi_modelos.csv"),
                 header_modelos, modelos_linhas)

    # -------- bi_checks.csv --------
    checks_linhas = []
    for row in sorted(MAPA_API.keys()):
        if row in CHECKS_BINARIOS:
            tipo = "binario"
        elif row in CHECKS_MAIOR_MELHOR:
            tipo = "maior_melhor"
        else:
            tipo = "menor_melhor"
        checks_linhas.append([
            row, MAPA_API[row], categorias[row],
            descricoes[row], pesos[row], tipo,
        ])
    escrever_csv(os.path.join(PASTA_BI, "bi_checks.csv"),
                 ["check_id", "chave_api", "categoria", "descricao",
                  "peso", "tipo"], checks_linhas)

    # -------- bi_metricas.csv --------
    metricas_linhas = []
    for arq in sorted(dados.keys()):
        metricas = dados[arq]
        if "_erro" in metricas:
            continue
        modelo_id = mapa_modelo_id[arq]
        for row, chave in MAPA_API.items():
            valor = metricas.get(chave)
            nota = metrica_para_nota(row, valor, escala)
            if nota is None:
                continue
            metricas_linhas.append([
                modelo_id, row, valor, nota,
                round(nota * pesos[row], 2),
            ])
    escrever_csv(os.path.join(PASTA_BI, "bi_metricas.csv"),
                 ["modelo_id", "check_id", "valor_bruto",
                  "nota", "nota_ponderada"], metricas_linhas)

    # -------- bi_resumo.csv --------
    resumo_linhas = []
    for arq in sorted(dados.keys()):
        metricas = dados[arq]
        if "_erro" in metricas:
            continue
        modelo_id = mapa_modelo_id[arq]
        notas = {}
        for row, chave in MAPA_API.items():
            nota = metrica_para_nota(row, metricas.get(chave), escala)
            if nota is not None:
                notas[row] = nota
        peso_total = sum(pesos[r] for r in notas)
        nota_final = (round(sum(notas[r] * pesos[r] for r in notas) / float(peso_total), 2)
                      if peso_total else 0)
        if nota_final >= 8.0:
            status = "Otimo"
        elif nota_final >= 6.0:
            status = "Bom"
        elif nota_final >= 4.0:
            status = "Regular"
        else:
            status = "Critico"
        resumo_linhas.append([modelo_id, nota_final, status, hoje])
    escrever_csv(os.path.join(PASTA_BI, "bi_resumo.csv"),
                 ["modelo_id", "nota_final", "status", "data_analise"],
                 resumo_linhas)

    # -------- sumario --------
    print(u"CSVs gerados em {}".format(PASTA_BI))
    print(u"  bi_modelos.csv   {:>4} linhas | {:>2} colunas".format(
        len(modelos_linhas), len(header_modelos)))
    print(u"  bi_checks.csv    {:>4} linhas".format(len(checks_linhas)))
    print(u"  bi_metricas.csv  {:>4} linhas".format(len(metricas_linhas)))
    print(u"  bi_resumo.csv    {:>4} linhas".format(len(resumo_linhas)))
    print(u"")
    print(u"Ranking:")
    for modelo_id, nota, status, _ in sorted(resumo_linhas, key=lambda x: -x[1]):
        arq = [a for a, m in mapa_modelo_id.items() if m == modelo_id][0]
        print(u"  {:>5}  {:10s}  {}".format("{:.2f}".format(nota), status, arq))


if __name__ == "__main__":
    main()