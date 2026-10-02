# -*- coding: utf-8 -*-
"""Consolida metricas_*.json em dashboard_bim.xlsx com calibracao relativa ao lote.

Calibracao:
  - Checks binarios (0/1): 0 -> 0, 1 -> 10
  - "Maior eh melhor": normaliza, mas 0 -> nota 0
  - "Menor eh melhor" (default): min-max invertido (melhor = 10, pior = 0)
  - Valor -1 ou None (falha): ignorado
"""
import os
import re
import sys
import glob
import json

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from bim_paths import PASTA_SAIDA, PLANILHA_AUDIT, XLSX_DASH
import openpyxl


# ============================================================
#  MAPEAMENTO: linha da aba ARC -> chave no JSON
# ============================================================
MAPA_API = {
    5:  "nomenclatura_arquivo",
    6:  "elementos_voados",
    8:  "dwg_refs",
    9:  "modelagem_local",
    10: "nivel_vista_fora_padrao",
    13: "folhas_fora_padrao",
    15: "carimbo_incompleto",
    16: "familias_generico",
    18: "tabelas_fora_padrao",
    23: "total_tags",
    25: "avisos_graves",
    75: "qtd_vinculos_rvt",           # NOVO
    76: "qtd_vinculos_cad",           # NOVO
    77: "qtd_vistas_sem_template",    # NOVO
}

# Checks que nao entram no dashboard (avaliacao humana)
ROWS_SUBJETIVOS = {
    7:  "Clash em ACC",
    11: "Organizacao de folhas",
    12: "Conteudo das folhas",
    14: "Preenchimento de Capa",
    17: "Fidelidade as categorias",
    19: "Identidade visual",
    20: "Nomenclatura geral de familias",
    24: "Configuracoes de publicacao",
}

# Checks excluidos por bug de API
CHECKS_EXCLUIDOS = {
    21: ("ponto_zero_ok",
         "ProjectLocation.Point nao mede ponto zero corretamente"),
    22: ("niveis_nao_fixados",
         "Level.IsPinned nao exposto no Revit 2020; fallback zerou o check"),
}

# Regex para recalcular nomenclatura do arquivo
RX_NOMENCLATURA = re.compile(
    r"^(WD[_\d]|\d+).*[-_](EX|EP|ARQ)[-_]?R\d+",
    re.IGNORECASE
)

# Classificacao dos checks
CHECKS_BINARIOS     = {5}
CHECKS_MAIOR_MELHOR = {23}


# ============================================================
#  PESOS (lidos da aba ARC)
# ============================================================
def carregar_pesos():
    wb = openpyxl.load_workbook(PLANILHA_AUDIT, data_only=True)
    ws = wb["ARC"]
    pesos, descricoes = {}, {}
    for row in MAPA_API:
        p = ws.cell(row=row, column=5).value
        d = ws.cell(row=row, column=2).value or ""
        pesos[row] = int(p) if isinstance(p, (int, float)) else 0
        descricoes[row] = str(d).strip()
    return pesos, descricoes


# ============================================================
#  CARREGAR JSONs
# ============================================================
def carregar_dados():
    dados = {}
    for jf in glob.glob(os.path.join(PASTA_SAIDA, "metricas_*.json")):
        try:
            with open(jf, encoding="utf-8") as f:
                dados.update(json.load(f))
        except Exception as e:
            print(u"[AVISO] Falha lendo {}: {}".format(jf, e))
    return dados


def recalcular_nomenclatura(dados):
    for arq, metricas in dados.items():
        nome = metricas.get("nome_arquivo") or arq
        metricas["nomenclatura_arquivo"] = int(bool(RX_NOMENCLATURA.match(nome)))
    return dados


# ============================================================
#  CALIBRACAO
# ============================================================
def calibrar_escala(dados):
    escala = {}
    for row in MAPA_API:
        chave = MAPA_API[row]
        valores = []
        for arq, metricas in dados.items():
            v = metricas.get(chave)
            if isinstance(v, (int, float)) and v >= 0:
                valores.append(v)
        if valores:
            escala[row] = {
                "min": min(valores),
                "max": max(valores),
                "valores": sorted(valores),
            }
    return escala


# ============================================================
#  REGUA
# ============================================================
def metrica_para_nota(row, valor, escala):
    if valor is None or valor < 0:
        return None
    if row not in escala:
        return None

    e = escala[row]
    vmin, vmax = e["min"], e["max"]

    if row in CHECKS_BINARIOS:
        return 10 if valor else 0

    if row in CHECKS_MAIOR_MELHOR:
        if valor == 0:
            return 0
        nao_zero = [x for x in e["valores"] if x > 0]
        if not nao_zero:
            return 0
        vmin_nz, vmax_nz = min(nao_zero), max(nao_zero)
        if vmax_nz == vmin_nz:
            return 10
        return round((valor - vmin_nz) / float(vmax_nz - vmin_nz) * 10, 2)

    # menor eh melhor (default)
    if vmax == vmin:
        return 10
    return round((1.0 - (valor - vmin) / float(vmax - vmin)) * 10, 2)


# ============================================================
#  MAIN
# ============================================================
def main():
    print(u"=" * 60)
    print(u"CONSOLIDADOR DE MATURIDADE BIM")
    print(u"=" * 60)
    print(u"Planilha audit : {}".format(PLANILHA_AUDIT))
    print(u"Pasta saida    : {}".format(PASTA_SAIDA))
    print(u"")

    if not os.path.isfile(PLANILHA_AUDIT):
        print(u"[ERRO] Planilha nao existe. Corrija PLANILHA_AUDIT em bim_paths.py")
        return
    if not os.path.isdir(PASTA_SAIDA):
        print(u"[ERRO] Pasta de saida nao existe: {}".format(PASTA_SAIDA))
        return

    try:
        pesos, descricoes = carregar_pesos()
    except Exception as e:
        print(u"[ERRO] Falha lendo pesos da planilha: {}".format(e))
        return
    print(u"Pesos carregados: {} checks".format(len(pesos)))

    # avisa se algum peso veio 0 (linha do Excel nao preenchida)
    zeros = [r for r in MAPA_API if pesos.get(r, 0) == 0]
    if zeros:
        print(u"[AVISO] Checks com peso=0 (nao serao contados): {}".format(zeros))

    dados = carregar_dados()
    if not dados:
        print(u"[ERRO] Nenhum metricas_*.json em {}".format(PASTA_SAIDA))
        return
    print(u"Modelos carregados: {}".format(len(dados)))

    dados = recalcular_nomenclatura(dados)
    n_ok = sum(1 for m in dados.values() if m.get("nomenclatura_arquivo"))
    print(u"Nomenclatura conforme regex: {}/{} modelos".format(n_ok, len(dados)))

    escala = calibrar_escala(dados)
    print(u"Escala calibrada para {} checks".format(len(escala)))
    print(u"")

    print(u"Valores brutos por check:")
    for row, chave in sorted(MAPA_API.items()):
        if row in escala:
            e = escala[row]
            print(u"  row {:2d} {:28s} min={:>6} max={:>6} n={}".format(
                row, chave[:28], e["min"], e["max"], len(e["valores"])))
        else:
            print(u"  row {:2d} {:28s} [sem dados]".format(row, chave[:28]))
    print(u"")

    # ============ ABA DETALHADO ============
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Detalhado"
    ws.append([u"Arquivo", u"Row", u"Descricao", u"Valor",
               u"Peso", u"Nota (0-10)", u"Nota x Peso"])

    notas_finais = {}
    for arq in sorted(dados.keys()):
        metricas = dados[arq]
        if "_erro" in metricas:
            notas_finais[arq] = None
            continue

        notas = {}
        for row, chave in MAPA_API.items():
            valor = metricas.get(chave)
            nota = metrica_para_nota(row, valor, escala)
            if nota is None:
                continue
            notas[row] = nota
            ws.append([arq, row, descricoes[row], valor, pesos[row],
                       nota, round(nota * pesos[row], 2)])

        peso_total = sum(pesos[r] for r in notas)
        nf = round(sum(notas[r] * pesos[r] for r in notas) / float(peso_total), 2) \
             if peso_total else 0
        notas_finais[arq] = nf
        ws.append([arq, "", u"NOTA FINAL", "", "", nf, ""])

    # ============ ABA RESUMO ============
    ws2 = wb.create_sheet("Resumo")
    ws2.append([u"Arquivo", u"Nota Final", u"Status"])
    for arq, nota in sorted(notas_finais.items(),
                            key=lambda kv: kv[1] or 0, reverse=True):
        if nota is None:
            status = u"N/A"
        elif nota >= 8.0:
            status = u"Otimo"
        elif nota >= 6.0:
            status = u"Bom"
        elif nota >= 4.0:
            status = u"Regular"
        else:
            status = u"Critico"
        ws2.append([arq, nota, status])

    # ============ ABA NAO MENSURAVEIS ============
    ws3 = wb.create_sheet(u"Nao mensuraveis")
    ws3.append([u"Row", u"Criterio", u"Motivo"])

    motivos_humanos = {
        7:  u"Requer Clash Detection (ACC/Navisworks)",
        11: u"Avaliacao subjetiva de leiaute",
        12: u"Requer inspecao humana de conteudo",
        14: u"Capa nao parametrizavel",
        17: u"Requer julgamento semantico de categorias",
        19: u"Identidade visual subjetiva",
        20: u"Nomenclatura requer julgamento humano",
        24: u"Dado externo (ACC)",
    }
    for row, desc in sorted(ROWS_SUBJETIVOS.items()):
        ws3.append([row, desc, motivos_humanos.get(row, u"")])
    for row, (chave, motivo) in sorted(CHECKS_EXCLUIDOS.items()):
        ws3.append([row, chave, motivo])

    # ============ SALVAR ============
    try:
        wb.save(XLSX_DASH)
    except Exception as e:
        print(u"[ERRO] Falha gravando XLSX: {}".format(e))
        return

    print(u"=" * 60)
    print(u"Excel gerado: {}".format(XLSX_DASH))
    print(u"")
    validos = [n for n in notas_finais.values() if n is not None]
    if validos:
        print(u"Estatisticas das notas finais:")
        print(u"  media   : {:.2f}".format(sum(validos) / len(validos)))
        print(u"  maximo  : {:.2f}".format(max(validos)))
        print(u"  minimo  : {:.2f}".format(min(validos)))
        if len(validos) > 1:
            s = sorted(validos)
            print(u"  mediana : {:.2f}".format(s[len(s) // 2]))
        print(u"")
        print(u"Ranking:")
        for arq, nota in sorted(notas_finais.items(),
                                key=lambda kv: kv[1] or 0, reverse=True):
            print(u"  {:>5}  {}".format(
                "{:.2f}".format(nota) if nota is not None else u"  N/A",
                arq))
    print(u"=" * 60)


if __name__ == "__main__":
    main()